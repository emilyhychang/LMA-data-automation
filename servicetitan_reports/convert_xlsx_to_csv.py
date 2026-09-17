#!/usr/bin/env python3
"""Convert one worksheet in an Excel workbook to CSV without changing the workbook."""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError as error:
    raise SystemExit(
        "This converter needs openpyxl. Install it once with: "
        "python3 -m pip install openpyxl"
    ) from error


def csv_value(value: object) -> object:
    """Keep CSV dates consistent and leave all other cell values unchanged."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def convert(source: Path, destination: Path, sheet_name: str | None) -> tuple[str, int]:
    """Export one worksheet and return its name and number of data rows."""
    if not source.is_file():
        raise FileNotFoundError(f"Excel file not found: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Choose an .xlsx or .xlsm file.")

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if sheet_name is None:
        worksheet = workbook.active
    elif sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
    else:
        choices = ", ".join(workbook.sheetnames)
        raise ValueError(f"Sheet '{sheet_name}' was not found. Available sheets: {choices}")

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        row_count = 0
        for row in worksheet.iter_rows(values_only=True):
            writer.writerow([csv_value(value) for value in row])
            row_count += 1
    workbook.close()
    return worksheet.title, max(row_count - 1, 0)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert an Excel worksheet to CSV without modifying the original workbook."
    )
    parser.add_argument("xlsx", type=Path, help="Workbook to convert (.xlsx or .xlsm)")
    parser.add_argument(
        "--sheet",
        help="Worksheet name. Defaults to the workbook's active sheet.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="CSV destination. Defaults to a CSV beside the workbook with the same name.",
    )
    args = parser.parse_args()
    destination = args.output or args.xlsx.with_suffix(".csv")
    try:
        sheet, rows = convert(args.xlsx, destination, args.sheet)
    except (FileNotFoundError, ValueError) as error:
        parser.error(str(error))
    print(f"Converted '{sheet}' to {destination} ({rows:,} data rows).")


if __name__ == "__main__":
    main()
